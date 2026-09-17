"""Extract exact dictionaries from the user's Excel."""
from openpyxl import load_workbook
from pathlib import Path
import json

PATH = Path(r"C:\Users\PZS0X\.minimax\v2\assets\2026\09\17\22-26-53-073-asset_20260917-222653-073_24913de2068a_a4c1617c-计划外维修汇总登记表.xlsx")
wb = load_workbook(PATH, data_only=True)

# Setting sheet: row 1 = headers, rows 2.. = values
# Cols: 单位 | 设备类型 | 报修内容 | 紧急程度 | 类型
setting = wb["Setting"]
units, equip_types, repair_contents, urgencies, type_names = [], [], [], [], []
for r in range(2, setting.max_row + 1):
    u = setting.cell(r, 1).value
    et = setting.cell(r, 2).value
    rc = setting.cell(r, 3).value
    ug = setting.cell(r, 4).value
    tn = setting.cell(r, 5).value
    if u: units.append(str(u).strip())
    if et: equip_types.append(str(et).strip())
    if rc: repair_contents.append(str(rc).strip())
    if ug: urgencies.append(str(ug).strip())
    if tn: type_names.append(str(tn).strip())

# Production sheet
prod = wb["Production"]
products = []
for r in range(2, prod.max_row + 1):
    v = prod.cell(r, 1).value
    if v: products.append(str(v).strip())

# Sample data row from the first sheet to understand which columns are used
print("SHEETS:", wb.sheetnames)
target = wb.sheetnames[0]
ws = wb[target]
print(f"Target: {target!r}")
sample = {}
for c in range(1, ws.max_column + 1):
    sample[ws.cell(1, c).value] = ws.cell(2, c).value

result = {
    "units": sorted(set(units)),
    "equip_types": sorted(set(equip_types)),
    "repair_contents": sorted(set(repair_contents)),
    "urgencies": sorted(set(urgencies)),
    "type_names": sorted(set(type_names)),
    "products": products,
    "sample_row": sample,
}
# Write seed-data.json (clean, no warnings) directly into the server tree.
out = {
    "units": sorted(set(units)),
    "equip_types": sorted(set(equip_types)),
    "repair_contents": sorted(set(repair_contents)),
    "urgencies": sorted(set(urgencies)),
    "type_names": sorted(set(type_names)),
    "products": products,
}
with open(r"D:\Github\Maintenance-Records\apps\server\src\cli\seed-data.json", "w", encoding="utf-8") as f:
    json.dump(out, f, ensure_ascii=False, indent=2)
print(f"Wrote seed-data.json: {len(out['units'])} units, {len(out['equip_types'])} equip types, {len(out['repair_contents'])} repair contents, {len(out['urgencies'])} urgencies, {len(out['type_names'])} types, {len(out['products'])} products")